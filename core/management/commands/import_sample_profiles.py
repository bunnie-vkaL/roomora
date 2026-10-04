"""Import the 100 fictional Hanoi profiles from the ROOMORA sample workbook."""

from datetime import date, datetime, time, timedelta
from decimal import Decimal

from django.conf import settings
from django.contrib.auth.models import User
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from core.constants import AREAS, QUESTIONS
from core.models import ImportedSampleProfile, LifestyleAnswers, LivingPreferences, Profile
from core.scoring import SCORING_VERSION


def source_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


SLEEP_TIMES = (
    (time(21, 45), time(22), time(22, 15)),
    (time(22, 45), time(23), time(23, 15)),
    (time(23, 45), time(0), time(0, 30)),
    (time(1, 15), time(1, 30), time(2)),
)
SURVEY_FALLBACKS = {
    "A1": 1, "A2": 0, "A3": 1, "B1": 1, "B2": 1, "C1": 1,
    "C2": 0, "C3": 1, "D1": 0, "D2": 0, "D3": 0,
    "D4": 1, "D5": 1, "E1": 1, "F1": 1, "F2": 0,
}


def simulated_living(item):
    """Make repeatable local demo preferences consistent with the survey answers."""

    answers = item["answers"]
    number = int(item["source_id"][2:])
    bedtime_group = answers.get("A1", number % len(SLEEP_TIMES))
    sleep_at = SLEEP_TIMES[bedtime_group][number % 3]
    sleep_minutes = sleep_at.hour * 60 + sleep_at.minute
    wake_minutes = sleep_minutes + 450 + (number % 3) * 30
    quiet_lead = (120, 90, 60, 30)[answers.get("D4", 1)]

    def clock(minutes):
        minutes %= 1440
        return time(minutes // 60, minutes % 60)

    needs = []
    if "B1" in answers:
        needs.append((
            "Ưu tiên không gian chung sạch, gọn mỗi ngày",
            "Muốn giữ không gian chung khá gọn",
            "Thoải mái với mức gọn gàng vừa phải",
            "Không quá khắt khe về sự ngăn nắp",
        )[answers["B1"]])
    if "C1" in answers:
        needs.append((
            "Cần tôn trọng phòng và đồ dùng riêng",
            "Thoải mái chia sẻ đồ dùng nếu hỏi trước",
            "Có thể dùng chung đồ đạc theo thỏa thuận",
        )[answers["C1"]])
    if answers.get("A3", 0) >= 2:
        needs.append("Cần góc làm việc tại nhà yên tĩnh")
    if answers.get("D2") == 2:
        needs.append("Cần nhà không có thú cưng")
    elif answers.get("D2") == 1:
        needs.append("Cần trao đổi trước nếu có thú cưng")
    elif answers.get("D3", 0) >= 1:
        needs.append("Muốn chỗ ở cho phép thú cưng")
    if answers.get("D1") == 0:
        needs.append("Ưu tiên không hút thuốc trong nhà")
    if answers.get("F1") == 0:
        needs.append("Không muốn khách ở qua đêm")
    needs.append(f"Mong người ở cùng cam kết tối thiểu {item['min_stay_months']} tháng")

    return {
        "move_in_from": item["move_in_from"],
        "move_in_until": item["move_in_from"] + timedelta(days=(14, 21, 28)[number % 3]),
        "needs": ". ".join(needs) + ".",
        "sleep_at": sleep_at,
        "wake_at": clock(wake_minutes),
        "quiet_from": clock(sleep_minutes - quiet_lead),
        "quiet_until": clock(wake_minutes),
        "total_monthly_budget": item["rent_max"] + 700_000 + (number % 3) * 150_000,
        "upfront_budget": item["rent_max"] * (2 + number % 2),
    }


def simulated_behavior(row, source_id):
    """Retain source metrics; fill absent values as explicitly simulated estimates."""

    number = int(source_id[2:])
    estimated = []

    def metric(index, key, fallback):
        if row[index] is None:
            estimated.append(key)
            return fallback
        return row[index]

    return {
        "payment_on_time_pct": int(metric(38, "payment_on_time_pct", 80 + number % 17)),
        "payment_periods": int(row[39] or 0),
        "reply_within_24h_pct": int(metric(40, "reply_within_24h_pct", 72 + number % 25)),
        "show_up_score": float(metric(41, "show_up_score", round(3.5 + (number % 14) / 10, 1))),
        "review_count": int(row[42] or 0),
        "review_average": float(metric(43, "review_average", round(3.7 + (number % 12) / 10, 1))),
        "verified_reports_90d": int(row[44] or 0),
        "estimated_fields": estimated,
    }


def parse_row(row, headers, row_number):
    source_id = str(row[0] or "").strip()
    if not source_id.startswith("RM") or not source_id[2:].isdigit():
        raise CommandError(f"Dòng {row_number}: ID hồ sơ không hợp lệ.")
    gender = {"Nam": "male", "Nữ": "female"}.get(row[2])
    if not gender:
        raise CommandError(f"Dòng {row_number}: giới tính không hợp lệ.")
    areas = [part.strip() for part in row[18].split(",")] if row[18] else list(AREAS)
    if not areas or any(area not in AREAS for area in areas):
        raise CommandError(f"Dòng {row_number}: H4 chứa khu vực không được hỗ trợ.")
    answers = {}
    estimated_answers = []
    for index, (key, _, _, options) in enumerate(QUESTIONS, 22):
        value = row[index]
        if value is None or value == "":
            answers[key] = SURVEY_FALLBACKS[key]
            estimated_answers.append(key)
            continue
        if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value or not 1 <= value <= len(options):
            raise CommandError(f"Dòng {row_number}: đáp án {key} không hợp lệ.")
        answers[key] = int(value) - 1
    try:
        rent_min = int(Decimal(str(row[16])) * 1_000_000)
        rent_max = int(Decimal(str(row[17])) * 1_000_000)
        age = int(row[4])
        birth_year = row[3].year
        move_in_from = row[19].date() if isinstance(row[19], datetime) else row[19]
        min_stay_months = int(row[21])
    except (AttributeError, TypeError, ValueError, ArithmeticError) as exc:
        raise CommandError(f"Dòng {row_number}: tuổi, ngày hoặc ngân sách không hợp lệ.") from exc
    if (not isinstance(move_in_from, date) or not 18 <= age <= 120 or not 0 < rent_min <= rent_max <= 100_000_000
            or not 1 <= min_stay_months <= 120):
        raise CommandError(f"Dòng {row_number}: tuổi, ngày hoặc ngân sách không hợp lệ.")
    source_data = {header: source_value(value) for header, value in zip(headers, row)}
    bio = " · ".join(str(value).strip() for value in (row[5], row[6]) if value)
    return {
        "source_id": source_id,
        "name": str(row[1]).strip(),
        "age": age,
        "gender": gender,
        "birth_year": birth_year,
        "bio": bio[:280],
        "areas": areas,
        "rent_min": rent_min,
        "rent_max": rent_max,
        "answers": answers,
        "estimated_answers": estimated_answers,
        "behavior_metrics": simulated_behavior(row, source_id),
        "move_in_from": move_in_from,
        "min_stay_months": min_stay_months,
        "source_data": source_data,
    }


class Command(BaseCommand):
    help = "Import fictional ROOMORA profiles from an XLSX workbook into the local development database."

    def add_arguments(self, parser):
        parser.add_argument("workbook", help="Path to ROOMORA_100_ho_so_mau_Ha_Noi.xlsx")
        parser.add_argument("--dry-run", action="store_true", help="Validate workbook without changing the database")

    def handle(self, *args, **options):
        if not settings.DEBUG:
            raise CommandError("Chỉ nhập hồ sơ mẫu khi DEBUG=True.")
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError("Cần cài openpyxl từ requirements.txt để đọc XLSX.") from exc
        try:
            workbook = load_workbook(options["workbook"], read_only=True, data_only=True)
            sheet = workbook["Hồ sơ"]
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows)
            if len(headers) != 48 or headers[0] != "ID" or headers[18] != "H4 Khu vực muốn tìm (quận)":
                raise CommandError("Sheet Hồ sơ không đúng cấu trúc 48 cột mong đợi.")
            parsed = [parse_row(row, headers, number) for number, row in enumerate(rows, 2) if row[0]]
            workbook.close()
        except (OSError, KeyError, StopIteration) as exc:
            raise CommandError(f"Không đọc được sheet Hồ sơ: {exc}") from exc
        ids = [item["source_id"] for item in parsed]
        if len(set(ids)) != len(ids):
            raise CommandError("Workbook có ID hồ sơ trùng lặp.")
        estimated_answer_count = sum(len(item["estimated_answers"]) for item in parsed)
        estimated_behavior_count = sum(len(item["behavior_metrics"]["estimated_fields"]) for item in parsed)
        if options["dry_run"]:
            self.stdout.write(
                f"Hợp lệ: {len(parsed)} hồ sơ đủ 16 câu; điền ước tính {estimated_answer_count} đáp án "
                f"và {estimated_behavior_count} chỉ số hành vi còn trống; chưa ghi database."
            )
            return
        with transaction.atomic():
            for item in parsed:
                source_id = item["source_id"]
                username = f"sample-{source_id.lower()}@roomora.local"
                user = User.objects.filter(username=username).first()
                if user and (not hasattr(user, "profile") or not hasattr(user.profile, "sample_source")
                             or user.profile.sample_source.source_id != source_id):
                    raise CommandError(f"Tài khoản {username} đã tồn tại và không thuộc bộ dữ liệu này.")
                if user is None:
                    user = User(username=username)
                    user.set_unusable_password()
                    user.save()
                profile, _ = Profile.objects.get_or_create(user=user, defaults={
                    "name": item["name"], "age": item["age"], "gender": item["gender"],
                    "birth_year": item["birth_year"], "areas": item["areas"],
                    "rent_min": item["rent_min"], "rent_max": item["rent_max"],
                    "contact_type": "phone", "contact_value": "Dữ liệu mẫu",
                })
                for field in ("name", "age", "gender", "birth_year", "bio", "areas", "rent_min", "rent_max"):
                    setattr(profile, field, item[field])
                profile.is_synthetic = True
                profile.is_published = True
                profile.questionnaire_version = SCORING_VERSION
                profile.full_clean()
                profile.save()
                LifestyleAnswers.objects.update_or_create(profile=profile, defaults={"values": item["answers"]})
                LivingPreferences.objects.update_or_create(profile=profile, defaults=simulated_living(item))
                ImportedSampleProfile.objects.update_or_create(profile=profile, defaults={
                    "source_id": source_id, "source_data": item["source_data"],
                    "estimated_answers": item["estimated_answers"],
                    "behavior_metrics": item["behavior_metrics"],
                })
        self.stdout.write(self.style.SUCCESS(
            f"Đã nhập/cập nhật {len(parsed)} hồ sơ và nhu cầu/nhịp sống giả lập; "
            f"cả {len(parsed)} hồ sơ đủ khảo sát để tham gia gợi ý."
        ))
